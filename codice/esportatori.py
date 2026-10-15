#!/usr/bin/env python3
"""
Esportatori: producono i file di consegna (.xlsx settimanale, riepilogo
annuale, bollettini Word) a partire da un oggetto sessione.Settimana.

Non modificare la struttura dei fogli 'riepilogo_settimana' e 'dati_grezzi'
(vedi CLAUDE.md): le righe sono fisse (pollini 6-52, spore 58-69).
"""

import re
from datetime import datetime, timedelta
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.styles.differential import DifferentialStyle
    from openpyxl.formatting.formatting import ConditionalFormattingList
    from openpyxl.formatting.rule import Rule
    from openpyxl.utils import get_column_letter
except ImportError:
    openpyxl = None

try:
    import docx as _docx_module
except ImportError:
    _docx_module = None

import dominio
import percorsi

TEMPLATE_FILE = percorsi.BUNDLE_DIR / "Polline_Template_Settimanale.xlsx"
SOGLIE_FILE_NOME = "concentrazioni_polliniche.xlsx"

# ── Stili openpyxl riutilizzati in piu' funzioni ──
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
) if openpyxl else None
FILL_GIALLO = PatternFill("solid", fgColor="FFE699") if openpyxl else None
FILL_VERDE = PatternFill("solid", fgColor="92D050") if openpyxl else None
FILL_VERDE_CHIARO = PatternFill("solid", fgColor="C5E0B4") if openpyxl else None
FILL_BLU = PatternFill("solid", fgColor="4472C4") if openpyxl else None
FONT_BOLD = Font(bold=True) if openpyxl else None
FONT_BOLD_BIG = Font(bold=True, size=14) if openpyxl else None
FONT_BIANCO_BOLD = Font(color="FFFFFF", bold=True) if openpyxl else None
ALIGN_CENTER = Alignment(horizontal="center") if openpyxl else None
ALIGN_CENTER_WRAP = Alignment(horizontal="center", wrap_text=True) if openpyxl else None


# ============================================================
# Soglie di concentrazione — unica fonte: concentrazioni_polliniche.xlsx
# ============================================================
def carica_soglie(output_dir=None):
    """Carica le soglie di concentrazione dal file esterno
    concentrazioni_polliniche.xlsx (unica fonte). Cerca in output_dir, poi accanto allo script,
    poi nel bundle. Ritorna dict {famiglia: (max_assente, max_bassa, max_media)}
    o None se il file non e' stato trovato."""
    candidati = []
    if output_dir is not None:
        candidati.append(Path(output_dir) / SOGLIE_FILE_NOME)
    candidati += [
        percorsi.SCRIPT_DIR / SOGLIE_FILE_NOME,
        percorsi.BUNDLE_DIR / SOGLIE_FILE_NOME,
    ]
    filepath = next((c for c in candidati if c.exists()), None)
    if filepath is None:
        return None

    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    try:
        ws = wb.active
        righe = (
            (row[0].value, row[1].value, row[2].value, row[3].value)
            for row in ws.iter_rows(min_row=3, max_col=5)
        )
        return dominio.parse_soglie(righe)
    finally:
        wb.close()


# ============================================================
# Esportazione del file settimanale (.xlsx)
# ============================================================
def compila_intestazione(ws, lunedi):
    """Compila i metadati della settimana nel foglio riepilogo (riga 3)."""
    domenica = lunedi + timedelta(days=6)
    mese_nome = dominio.MESI_ITA[lunedi.month - 1]
    anno = lunedi.year
    fmt = "%d-%m-%Y"

    ws["H3"] = mese_nome
    ws["J3"] = lunedi
    ws["J3"].number_format = "DD-MM-YYYY"
    ws["K3"] = domenica
    ws["K3"].number_format = "DD-MM-YYYY"
    ws["M3"] = anno

    ws["T3"] = mese_nome
    ws["V3"] = lunedi.strftime(fmt)
    ws["W3"] = domenica.strftime(fmt)
    ws["Y3"] = anno


# Bollettino nel foglio riepilogo_settimana: la riga _BOLL_PRIMA_RIGA + i
# corrisponde a dominio.BOLLETTINO_RIGHE[i] (righe 76-110). Due copie affiancate,
# ciascuna con 7 giorni + media: E-L e Q-X.
_BOLL_PRIMA_RIGA = 76
_BOLL_LATI = (("E", "L"), ("Q", "X"))


def _regole_livelli(soglia_tuple):
    """Le quattro regole di colore di una riga, in ordine di priorita'. Stessa
    semantica di dominio.livello(): ogni soglia e' il massimo del livello sotto.
    Il riempimento ha sia fgColor sia bgColor: i formati condizionali di Excel
    leggono il secondo, LibreOffice il primo."""
    max_ass, max_bas, max_med = soglia_tuple
    regole = []
    for livello, operatore, valore in (
        ("alta", "greaterThan", max_med),
        ("media", "greaterThan", max_bas),
        ("bassa", "greaterThan", max_ass),
        ("assente", "greaterThanOrEqual", 0),
    ):
        colore = dominio.LIVELLO_COLORE_WORD[livello]
        stile = DifferentialStyle(
            fill=PatternFill(fill_type="solid", start_color=colore, end_color=colore),
            font=Font(color="FFFFFF") if livello == "alta" else None)
        regole.append(Rule(type="cellIs", operator=operatore, formula=[repr(valore)],
                           dxf=stile, stopIfTrue=True))
    return regole


def colora_bollettino(ws, soglie):
    """Formattazione condizionale delle due tabelle del bollettino, con le soglie di
    concentrazioni_polliniche.xlsx (la stessa fonte della GUI e del Word). Sostituisce
    le regole del template, che avevano soglie scritte a mano."""
    ws.conditional_formatting = ConditionalFormattingList()
    for i, (_ita, _eng, _codici, famiglia) in enumerate(dominio.BOLLETTINO_RIGHE):
        riga = _BOLL_PRIMA_RIGA + i
        soglia_tuple = dominio.soglia_famiglia(soglie, famiglia)
        for da, a in _BOLL_LATI:
            intervallo = f"{da}{riga}:{a}{riga}"
            for regola in _regole_livelli(soglia_tuple):
                ws.conditional_formatting.add(intervallo, regola)


def esporta_xlsx(settimana, path, soglie=None):
    """Scrive il file .xlsx settimanale (template + conteggi + log) su path.
    soglie: dict da carica_soglie(); se None viene caricato."""
    path = Path(path)
    if soglie is None:
        soglie = carica_soglie(path.parent) or {}
    wb = openpyxl.load_workbook(TEMPLATE_FILE)
    ws = wb["riepilogo_settimana"]
    ws_log = wb["dati_grezzi"]

    compila_intestazione(ws, settimana.lunedi)
    ws["Q3"] = settimana.fattore

    for codice in dominio.TUTTI_CODICI:
        row = dominio.codice_to_row(codice)
        for g in range(1, 8):
            col = dominio.giorno_to_col(g)
            v = settimana.conteggi[codice][g - 1]
            ws.cell(row=row, column=col, value=v if v else None)

    for i, riga in enumerate(settimana.log, start=2):
        ws_log.cell(row=i, column=1, value=riga["data"])
        ws_log.cell(row=i, column=2, value=riga["codice"])
        ws_log.cell(row=i, column=3, value=riga["specie"])
        ws_log.cell(row=i, column=4, value=riga["ora"])
        if riga.get("nota"):
            ws_log.cell(row=i, column=5, value=riga["nota"])

    colora_bollettino(ws, soglie)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()
    return path


# ============================================================
# Riepilogo annuale
# ============================================================
POLLINI_CODICI = dominio.POLLINI_CODICI
SPORE_CODICI = dominio.SPORE_CODICI

_ANN_POLL_START = 2                                           # col B
_ANN_SEP1 = _ANN_POLL_START + len(POLLINI_CODICI)             # 49
_ANN_SPORE_START = _ANN_SEP1 + 1                               # 50
_ANN_SPORE_END = _ANN_SPORE_START + len(SPORE_CODICI) - 1     # 61
_ANN_GAP = _ANN_SPORE_END + 1                                  # 62
_ANN_CONC_DATA = _ANN_GAP + 1                                   # 63
_ANN_CONC_POLL_START = _ANN_CONC_DATA + 1                      # 64
_ANN_CONC_SEP = _ANN_CONC_POLL_START + len(POLLINI_CODICI)     # 111
_ANN_CONC_SPORE_START = _ANN_CONC_SEP + 1                      # 112
_ANN_CONC_SPORE_END = _ANN_CONC_SPORE_START + len(SPORE_CODICI) - 1  # 123

# Fattore di conversione: una sola cella, nell'angolo superiore destro della
# tabella delle concentrazioni (come Q3 nel template settimanale). Le
# concentrazioni sono formule "conta x fattore" che puntano sempre a questa
# cella (riferimento assoluto) e alla conta della stessa riga.
_ANN_FATTORE_RIGA = 1
_ANN_FATTORE_COL = _ANN_CONC_SPORE_END                        # 123 (DS)
_FMT_CONC = "0.0;-0.0;;@"      # i valori nulli restano vuoti, come nel template
_FMT_CONTA = "0;-0;;@"


def _ann_col_grezzo(codice):
    return dominio.posizione_codice(codice, _ANN_POLL_START, _ANN_SPORE_START)


def _ann_col_conc(codice):
    return dominio.posizione_codice(codice, _ANN_CONC_POLL_START, _ANN_CONC_SPORE_START)


def _fill_specie(codice):
    """Riempimento delle specie di interesse nei riepiloghi (None per le altre)."""
    if codice in dominio.ANNUALE_VERDE_CODICI:
        return FILL_VERDE
    if codice in dominio.ANNUALE_VERDE_CHIARO_CODICI:
        return FILL_VERDE_CHIARO
    return None


def _evidenzia_specie(cell, codice, grassetto=True):
    """Riempimento (ed eventuale grassetto) di una cella di una specie di interesse."""
    fill = _fill_specie(codice)
    if fill is not None:
        cell.fill = fill
        if grassetto:
            cell.font = FONT_BOLD


def _rif_fattore():
    """Riferimento assoluto alla cella del fattore nel foglio dati (es. $DS$1)."""
    return f"${get_column_letter(_ANN_FATTORE_COL)}${_ANN_FATTORE_RIGA}"


def _formula_concentrazione(col_grezzo, riga):
    """Formula conta x fattore per la riga: vuota se non c'e' conta."""
    ref = f"{get_column_letter(col_grezzo)}{riga}"
    return f'=IF({ref}="","",{ref}*{_rif_fattore()})'


def _rif_foglio(titolo):
    return "'" + titolo.replace("'", "''") + "'"


def _scrivi_fattore_annuale(ws, fattore):
    etichetta = ws.cell(row=_ANN_FATTORE_RIGA, column=_ANN_FATTORE_COL - 1, value="Fatt.conv")
    etichetta.font = FONT_BOLD
    etichetta.alignment = Alignment(horizontal="right")
    cella = ws.cell(row=_ANN_FATTORE_RIGA, column=_ANN_FATTORE_COL, value=fattore)
    cella.font = FONT_BOLD
    cella.fill = FILL_GIALLO
    cella.border = THIN_BORDER
    cella.alignment = ALIGN_CENTER
    ws.column_dimensions[get_column_letter(_ANN_FATTORE_COL)].width = 5


def _leggi_fattore_annuale(ws):
    """Fattore scritto nel foglio dati, o None se manca (file creato prima
    che il riepilogo annuale avesse la cella del fattore)."""
    v = ws.cell(row=_ANN_FATTORE_RIGA, column=_ANN_FATTORE_COL).value
    if isinstance(v, (int, float)) and not isinstance(v, bool) and v > 0:
        return float(v)
    return None


def _deduci_fattore_file_vecchio(ws):
    """In un file senza cella del fattore le concentrazioni erano numeri
    fissi: il fattore usato si ricava dal primo rapporto concentrazione/conta
    (None se non ci sono dati)."""
    for riga in range(4, ws.max_row + 1):
        for codice in POLLINI_CODICI + SPORE_CODICI:
            conta = ws.cell(row=riga, column=_ann_col_grezzo(codice)).value
            conc = ws.cell(row=riga, column=_ann_col_conc(codice)).value
            if (isinstance(conta, (int, float)) and conta > 0
                    and isinstance(conc, (int, float)) and conc > 0):
                return round(conc / conta, 4)
    return None


def _migra_a_formule(ws, ws_cal):
    """Converte in formule le concentrazioni di un file creato quando erano
    numeri fissi. Nel Calendario si convertono le colonne la cui data compare
    una sola volta nel foglio dati (con date doppie la corrispondenza non e'
    certa e il valore resta com'e')."""
    righe_per_data = {}
    for riga in range(4, ws.max_row + 1):
        data = ws.cell(row=riga, column=1).value
        if data is None:
            continue
        righe_per_data.setdefault(str(data).strip(), []).append(riga)
        for codice in POLLINI_CODICI + SPORE_CODICI:
            ws.cell(row=riga, column=_ann_col_conc(codice)).value = \
                _formula_concentrazione(_ann_col_grezzo(codice), riga)
    for col in range(_CAL_COL_DATE_START, ws_cal.max_column + 1):
        data = ws_cal.cell(row=_CAL_HEADER_ROW, column=col).value
        righe = righe_per_data.get(str(data).strip()) if data else None
        if righe and len(righe) == 1:
            _scrivi_formule_calendario(ws_cal, col, ws.title, righe[0])


def crea_intestazione_annuale(ws, anno, fattore=dominio.FATTORE_DEFAULT):
    cell = ws.cell(row=1, column=1, value=f"RIEPILOGO ANNUALE {anno}")
    cell.font = FONT_BOLD_BIG
    cell.alignment = ALIGN_CENTER
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=_ANN_SPORE_END)

    cell = ws.cell(row=2, column=1, value="CONTA GREZZA")
    cell.font = FONT_BOLD
    cell.alignment = ALIGN_CENTER
    cell.fill = FILL_GIALLO
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=_ANN_SPORE_END)

    cell = ws.cell(row=2, column=_ANN_CONC_DATA, value="CONCENTRAZIONI (p/m3)")
    cell.font = FONT_BOLD
    cell.alignment = ALIGN_CENTER
    cell.fill = FILL_GIALLO
    ws.merge_cells(start_row=2, start_column=_ANN_CONC_DATA,
                   end_row=2, end_column=_ANN_CONC_SPORE_END)

    align_rotated = Alignment(horizontal="center", text_rotation=90)

    def _scrivi_intestazione_colonna(col, codice):
        nome = dominio.CODICI_SPECIE[codice]
        cell = ws.cell(row=3, column=col, value=nome)
        cell.font = FONT_BOLD if codice in dominio.ANNUALE_BOLD_CODICI else Font()
        cell.border = THIN_BORDER
        cell.alignment = align_rotated
        fill = _fill_specie(codice)
        cell.fill = fill if fill is not None else FILL_GIALLO

    cell = ws.cell(row=3, column=1, value="Data")
    cell.font = FONT_BOLD
    cell.fill = FILL_GIALLO
    cell.border = THIN_BORDER
    cell.alignment = ALIGN_CENTER

    for i, codice in enumerate(POLLINI_CODICI):
        _scrivi_intestazione_colonna(_ANN_POLL_START + i, codice)

    cell = ws.cell(row=3, column=_ANN_SEP1, value="||")
    cell.font = FONT_BOLD
    cell.border = THIN_BORDER
    cell.alignment = ALIGN_CENTER

    for i, codice in enumerate(SPORE_CODICI):
        _scrivi_intestazione_colonna(_ANN_SPORE_START + i, codice)

    cell = ws.cell(row=3, column=_ANN_CONC_DATA, value="Data")
    cell.font = FONT_BOLD
    cell.fill = FILL_GIALLO
    cell.border = THIN_BORDER
    cell.alignment = ALIGN_CENTER

    for i, codice in enumerate(POLLINI_CODICI):
        _scrivi_intestazione_colonna(_ANN_CONC_POLL_START + i, codice)

    cell = ws.cell(row=3, column=_ANN_CONC_SEP, value="||")
    cell.font = FONT_BOLD
    cell.border = THIN_BORDER
    cell.alignment = ALIGN_CENTER

    for i, codice in enumerate(SPORE_CODICI):
        _scrivi_intestazione_colonna(_ANN_CONC_SPORE_START + i, codice)

    ws.column_dimensions["A"].width = 11
    for i in range(len(POLLINI_CODICI)):
        ws.column_dimensions[get_column_letter(_ANN_POLL_START + i)].width = 4
    ws.column_dimensions[get_column_letter(_ANN_SEP1)].width = 2
    for i in range(len(SPORE_CODICI)):
        ws.column_dimensions[get_column_letter(_ANN_SPORE_START + i)].width = 4
    ws.column_dimensions[get_column_letter(_ANN_GAP)].width = 3
    ws.column_dimensions[get_column_letter(_ANN_CONC_DATA)].width = 11
    for i in range(len(POLLINI_CODICI)):
        ws.column_dimensions[get_column_letter(_ANN_CONC_POLL_START + i)].width = 4
    ws.column_dimensions[get_column_letter(_ANN_CONC_SEP)].width = 2
    for i in range(len(SPORE_CODICI)):
        ws.column_dimensions[get_column_letter(_ANN_CONC_SPORE_START + i)].width = 4

    ws.row_dimensions[3].height = 80
    ws.freeze_panes = "B4"
    ws.auto_filter.ref = f"A3:{get_column_letter(_ANN_CONC_SPORE_END)}3"
    _scrivi_fattore_annuale(ws, fattore)


def trova_riga_per_data(ws, data_str):
    for row in range(4, ws.max_row + 1):
        val = ws.cell(row=row, column=1).value
        if val and str(val).strip() == data_str:
            return row
    return None


def _prossima_riga_annuale(ws):
    for row in range(4, ws.max_row + 2):
        if ws.cell(row=row, column=1).value is None:
            return row
    return ws.max_row + 1


def scrivi_riga_annuale(ws, riga, data_str, dati, modo):
    """modo: 'nuovo'/'sovrascrivi' = scrive i valori; 'somma' = aggiunge agli
    esistenti. Le concentrazioni sono formule (conta x cella del fattore):
    si riscrivono sempre, quindi 'somma' agisce solo sulla conta."""
    cell = ws.cell(row=riga, column=1, value=data_str)
    cell.border = THIN_BORDER

    cell = ws.cell(row=riga, column=_ANN_CONC_DATA, value=data_str)
    cell.border = THIN_BORDER

    for codice in POLLINI_CODICI + SPORE_CODICI:
        val_nuovo = dati.get(codice, 0)
        col_grezzo = _ann_col_grezzo(codice)
        col_conc = _ann_col_conc(codice)
        if col_grezzo is None:
            continue

        if modo == "somma" and val_nuovo > 0:
            val_esistente = ws.cell(row=riga, column=col_grezzo).value
            if isinstance(val_esistente, (int, float)):
                val_nuovo = int(val_esistente) + val_nuovo

        cell_g = ws.cell(row=riga, column=col_grezzo)
        cell_g.value = val_nuovo if val_nuovo > 0 else None
        cell_g.border = THIN_BORDER
        cell_g.alignment = ALIGN_CENTER
        _evidenzia_specie(cell_g, codice)

        cell_c = ws.cell(row=riga, column=col_conc)
        cell_c.value = _formula_concentrazione(col_grezzo, riga)
        cell_c.border = THIN_BORDER
        cell_c.alignment = ALIGN_CENTER
        cell_c.number_format = _FMT_CONC
        _evidenzia_specie(cell_c, codice)

    for sep_col in (_ANN_SEP1, _ANN_CONC_SEP):
        cell = ws.cell(row=riga, column=sep_col, value="||")
        cell.border = THIN_BORDER
        cell.alignment = ALIGN_CENTER
    return riga


def _nome_foglio_settimana(lunedi):
    return f"W{lunedi.isocalendar()[1]:02d}"


def _posizione_foglio_settimana(wb_ann, nome_foglio):
    nomi = wb_ann.sheetnames
    num_fissi = sum(1 for n in nomi if not re.match(r"^W\d+$", n))
    nuovo_num = int(nome_foglio[1:])
    pos = num_fissi
    for n in nomi:
        if re.match(r"^W\d+$", n) and int(n[1:]) < nuovo_num:
            pos += 1
    return pos


def crea_foglio_settimana_annuale(wb_ann, conteggi, lunedi, fattore):
    """Crea o sovrascrive il foglio settimanale (es. 'W08') nel riepilogo annuale."""
    nome_foglio = _nome_foglio_settimana(lunedi)
    domenica = lunedi + timedelta(days=6)
    settimana_num = lunedi.isocalendar()[1]

    if nome_foglio in wb_ann.sheetnames:
        idx = wb_ann.sheetnames.index(nome_foglio)
        del wb_ann[nome_foglio]
        ws_s = wb_ann.create_sheet(nome_foglio, idx)
    else:
        pos = _posizione_foglio_settimana(wb_ann, nome_foglio)
        ws_s = wb_ann.create_sheet(nome_foglio, pos)

    titolo = (f"SETTIMANA {settimana_num}  -  "
              f"{lunedi.strftime('%d/%m/%Y')} / {domenica.strftime('%d/%m/%Y')}")
    c = ws_s.cell(row=1, column=1, value=titolo)
    c.font = FONT_BOLD_BIG
    ws_s.merge_cells(start_row=1, start_column=1, end_row=1, end_column=9)

    giorni_hdr = []
    for g in range(7):
        dt = lunedi + timedelta(days=g)
        giorni_hdr.append(f"{dominio.GIORNI_ABBREV[g]}\n{dt.strftime('%d/%m')}")

    def _scrivi_header(riga, label_ultima, fill_hdr, font_hdr):
        c = ws_s.cell(row=riga, column=1, value="Specie")
        c.font = font_hdr; c.fill = fill_hdr; c.border = THIN_BORDER; c.alignment = ALIGN_CENTER
        for g, hdr in enumerate(giorni_hdr):
            c = ws_s.cell(row=riga, column=2 + g, value=hdr)
            c.font = font_hdr; c.fill = fill_hdr; c.border = THIN_BORDER
            c.alignment = ALIGN_CENTER_WRAP
        c = ws_s.cell(row=riga, column=9, value=label_ultima)
        c.font = font_hdr; c.fill = fill_hdr; c.border = THIN_BORDER; c.alignment = ALIGN_CENTER

    def _scrivi_specie(riga, codice, vals, valore_finale, fmt="0", evidenzia=None):
        """``vals``: numeri o formule; ``evidenzia``: per ogni giorno, se la
        cella va in grassetto (di norma quando c'e' un valore)."""
        if evidenzia is None:
            evidenzia = [bool(v) for v in vals]
        specie = dominio.CODICI_SPECIE[codice]
        c = ws_s.cell(row=riga, column=1, value=specie)
        c.border = THIN_BORDER
        _evidenzia_specie(c, codice)
        for g, v in enumerate(vals):
            c = ws_s.cell(row=riga, column=2 + g, value=(v if v else None))
            c.border = THIN_BORDER; c.alignment = ALIGN_CENTER; c.number_format = fmt
            _evidenzia_specie(c, codice, grassetto=evidenzia[g])
        c = ws_s.cell(row=riga, column=9, value=(valore_finale if valore_finale else None))
        c.border = THIN_BORDER; c.alignment = ALIGN_CENTER; c.number_format = fmt

    def _scrivi_separatore(riga):
        for col in range(1, 10):
            ws_s.cell(row=riga, column=col).border = THIN_BORDER

    def _scrivi_doy(riga, fill_hdr, font_hdr):
        c = ws_s.cell(row=riga, column=1, value="G. anno")
        c.border = THIN_BORDER; c.fill = fill_hdr; c.font = font_hdr; c.alignment = ALIGN_CENTER
        for g in range(7):
            dt_g = lunedi + timedelta(days=g)
            doy = dt_g.timetuple().tm_yday
            c = ws_s.cell(row=riga, column=2 + g, value=doy)
            c.border = THIN_BORDER; c.fill = fill_hdr; c.font = font_hdr; c.alignment = ALIGN_CENTER
        c = ws_s.cell(row=riga, column=9)
        c.border = THIN_BORDER; c.fill = fill_hdr

    r = 2
    c = ws_s.cell(row=r, column=1, value="CONTA GREZZA")
    c.font = FONT_BOLD; c.fill = FILL_GIALLO; c.alignment = ALIGN_CENTER
    ws_s.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)

    r = 3
    _scrivi_header(r, "Totale sett.", FILL_GIALLO, FONT_BOLD)
    r = 4
    _scrivi_doy(r, FILL_GIALLO, FONT_BOLD)

    r = 5
    for codice in POLLINI_CODICI:
        vals = conteggi.get(codice, [0] * 7)
        _scrivi_specie(r, codice, vals, sum(vals))
        r += 1
    _scrivi_separatore(r); r += 1
    for codice in SPORE_CODICI:
        vals = conteggi.get(codice, [0] * 7)
        _scrivi_specie(r, codice, vals, sum(vals))
        r += 1

    r += 1
    c = ws_s.cell(row=r, column=1, value="CONCENTRAZIONI (p/m3)")
    c.font = FONT_BIANCO_BOLD; c.fill = FILL_BLU; c.alignment = ALIGN_CENTER
    ws_s.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
    # Fattore di conversione di questa settimana, accanto al titolo: le
    # concentrazioni sotto sono formule "conta x fattore" che puntano a questa
    # cella (stesso schema del template settimanale).
    c = ws_s.cell(row=r, column=11, value="Fatt.conv")
    c.font = FONT_BOLD; c.alignment = Alignment(horizontal="right")
    c = ws_s.cell(row=r, column=12, value=fattore)
    c.font = FONT_BOLD; c.fill = FILL_GIALLO; c.border = THIN_BORDER; c.alignment = ALIGN_CENTER
    rif_fattore = f"$L${r}"

    r += 1
    _scrivi_header(r, "Media sett.", FILL_BLU, FONT_BIANCO_BOLD)
    r += 1
    _scrivi_doy(r, FILL_BLU, FONT_BIANCO_BOLD)

    r += 1
    scarto = r - 5          # riga delle concentrazioni = riga della conta + scarto

    def _scrivi_concentrazioni(codice):
        vals_raw = conteggi.get(codice, [0] * 7)
        riga_conta = r - scarto
        formule = []
        for g in range(7):
            ref = f"{get_column_letter(2 + g)}{riga_conta}"
            formule.append(f'=IF({ref}="","",{ref}*{rif_fattore})')
        _scrivi_specie(r, codice, formule, f"=SUM(B{r}:H{r})/7", fmt=_FMT_CONC,
                       evidenzia=[bool(v) for v in vals_raw])

    for codice in POLLINI_CODICI:
        _scrivi_concentrazioni(codice)
        r += 1
    _scrivi_separatore(r); r += 1
    for codice in SPORE_CODICI:
        _scrivi_concentrazioni(codice)
        r += 1

    ws_s.column_dimensions["A"].width = 28
    for g in range(7):
        ws_s.column_dimensions[get_column_letter(2 + g)].width = 8
    ws_s.column_dimensions[get_column_letter(9)].width = 9
    ws_s.column_dimensions["K"].width = 10
    ws_s.column_dimensions["L"].width = 7
    ws_s.row_dimensions[3].height = 32

    return ws_s


# ============================================================
# Foglio Calendario (trasposto: specie in righe, date in colonne)
# ============================================================
_CAL_HEADER_ROW = 3
_CAL_DOY_ROW = 4
_CAL_DATA_START_ROW = 5
_CAL_COL_DATE_START = 2
_CAL_SEP_ROW = _CAL_DATA_START_ROW + len(POLLINI_CODICI)  # 52


def _cal_row_for_codice(codice):
    return dominio.posizione_codice(codice, _CAL_DATA_START_ROW, _CAL_SEP_ROW + 1)


def crea_intestazione_calendario(ws, anno):
    cell = ws.cell(row=1, column=1, value=f"CALENDARIO POLLINICO {anno}")
    cell.font = FONT_BOLD_BIG
    cell.alignment = ALIGN_CENTER

    cell = ws.cell(row=2, column=1, value="CONCENTRAZIONI (p/m3)")
    cell.font = FONT_BOLD
    cell.alignment = ALIGN_CENTER
    cell.fill = FILL_GIALLO

    cell = ws.cell(row=_CAL_HEADER_ROW, column=1, value="Specie")
    cell.font = FONT_BOLD
    cell.fill = FILL_GIALLO
    cell.border = THIN_BORDER
    cell.alignment = ALIGN_CENTER

    cell = ws.cell(row=_CAL_DOY_ROW, column=1, value="G. anno")
    cell.font = FONT_BOLD
    cell.fill = FILL_GIALLO
    cell.border = THIN_BORDER
    cell.alignment = ALIGN_CENTER

    for i, codice in enumerate(POLLINI_CODICI):
        row = _CAL_DATA_START_ROW + i
        cell = ws.cell(row=row, column=1, value=dominio.CODICI_SPECIE[codice])
        cell.font = FONT_BOLD if codice in dominio.ANNUALE_BOLD_CODICI else Font()
        cell.border = THIN_BORDER
        fill = _fill_specie(codice)
        if fill is not None:
            cell.fill = fill

    cell = ws.cell(row=_CAL_SEP_ROW, column=1, value="||")
    cell.border = THIN_BORDER
    cell.alignment = ALIGN_CENTER

    for i, codice in enumerate(SPORE_CODICI):
        row = _CAL_SEP_ROW + 1 + i
        cell = ws.cell(row=row, column=1, value=dominio.CODICI_SPECIE[codice])
        cell.font = FONT_BOLD if codice in dominio.ANNUALE_BOLD_CODICI else Font()
        cell.border = THIN_BORDER
        fill = _fill_specie(codice)
        if fill is not None:
            cell.fill = fill

    ws.column_dimensions["A"].width = 26
    ws.row_dimensions[_CAL_HEADER_ROW].height = 60
    ws.row_dimensions[_CAL_DOY_ROW].height = 16
    ws.freeze_panes = "B5"


def trova_colonna_per_data_calendario(ws, data_str):
    for col in range(_CAL_COL_DATE_START, ws.max_column + 1):
        val = ws.cell(row=_CAL_HEADER_ROW, column=col).value
        if val and str(val).strip() == data_str:
            return col
    return None


def _prossima_colonna_calendario(ws):
    for col in range(_CAL_COL_DATE_START, ws.max_column + 2):
        if ws.cell(row=_CAL_HEADER_ROW, column=col).value is None:
            return col
    return ws.max_column + 1


def _scrivi_formule_calendario(ws, col, titolo_dati, riga_dati):
    """Ogni cella del calendario rimanda alla concentrazione del foglio dati
    (stessa data, stessa specie): una sola fonte di calcolo."""
    foglio = _rif_foglio(titolo_dati)
    for codice in POLLINI_CODICI + SPORE_CODICI:
        row = _cal_row_for_codice(codice)
        if row is None:
            continue
        sorgente = f"{foglio}!{get_column_letter(_ann_col_conc(codice))}{riga_dati}"
        cell = ws.cell(row=row, column=col)
        cell.value = f"={sorgente}"
        cell.border = THIN_BORDER
        cell.alignment = ALIGN_CENTER
        cell.number_format = _FMT_CONC
        _evidenzia_specie(cell, codice)


def scrivi_colonna_calendario(ws, col, data_str, titolo_dati, riga_dati):
    """Colonna del calendario per ``data_str``: le concentrazioni sono
    riferimenti alla riga ``riga_dati`` del foglio ``titolo_dati``."""
    align_rotated = Alignment(horizontal="center", text_rotation=90)

    cell = ws.cell(row=_CAL_HEADER_ROW, column=col, value=data_str)
    cell.font = FONT_BOLD
    cell.fill = FILL_GIALLO
    cell.border = THIN_BORDER
    cell.alignment = align_rotated
    ws.column_dimensions[get_column_letter(col)].width = 5

    try:
        dt_col = datetime.strptime(data_str, "%d/%m/%Y")
        doy = dt_col.timetuple().tm_yday
        cell_doy = ws.cell(row=_CAL_DOY_ROW, column=col, value=doy)
        cell_doy.font = FONT_BOLD
        cell_doy.fill = FILL_GIALLO
        cell_doy.border = THIN_BORDER
        cell_doy.alignment = ALIGN_CENTER
    except ValueError:
        pass

    cell = ws.cell(row=_CAL_SEP_ROW, column=col, value="||")
    cell.border = THIN_BORDER
    cell.alignment = ALIGN_CENTER

    _scrivi_formule_calendario(ws, col, titolo_dati, riga_dati)


def raccogli_dati_giornalieri(conteggi):
    """{giorno_num: {codice: valore}} solo per giorni con dati > 0."""
    risultato = {}
    for giorno_num in range(1, 8):
        dati = {codice: conteggi[codice][giorno_num - 1] for codice in dominio.TUTTI_CODICI
                if conteggi[codice][giorno_num - 1] > 0}
        if dati:
            risultato[giorno_num] = dati
    return risultato


def esporta_riepilogo_annuale(settimana, cartella, scegli_duplicati=None, avvisa=None):
    """Esporta i dati settimanali nel file riepilogo annuale.

    scegli_duplicati: callback(data_str) -> 'a' (sovrascrivi) / 'b' (nuova riga)
    / 'c' (somma) / 'annulla' (lascia il giorno invariato, non lo esporta),
    chiamata solo se il giorno e' gia' presente. Se None e c'e' un duplicato,
    si sovrascrive senza chiedere.
    avvisa: callback(messaggio) chiamata se il file annuale usa gia' un
    fattore di conversione diverso da quello della settimana (il file tiene il
    proprio: l'unica cella del fattore vale per tutto il foglio).
    Ritorna (percorso, n_giorni_esportati) o (None, 0) se non c'e' nulla da esportare.
    """
    fattore = settimana.fattore
    anno = settimana.lunedi.year
    nome_file = f"Riepilogo_Annuale_{anno}.xlsx"
    percorso = Path(cartella) / nome_file

    dati_settimana = raccogli_dati_giornalieri(settimana.conteggi)
    if not dati_settimana:
        return None, 0

    migra = False
    if percorso.exists():
        wb_ann = openpyxl.load_workbook(percorso)
        ws = wb_ann.active
        fattore_file = _leggi_fattore_annuale(ws)
        if fattore_file is None:        # file creato prima della cella del fattore
            fattore_file = _deduci_fattore_file_vecchio(ws) or fattore
            _scrivi_fattore_annuale(ws, fattore_file)
            migra = True
    else:
        wb_ann = openpyxl.Workbook()
        ws = wb_ann.active
        ws.title = f"Dati {anno}"
        fattore_file = fattore
        crea_intestazione_annuale(ws, anno, fattore)

    if "Calendario" in wb_ann.sheetnames:
        ws_cal = wb_ann["Calendario"]
    else:
        ws_cal = wb_ann.create_sheet("Calendario")
        crea_intestazione_calendario(ws_cal, anno)

    if migra:
        _migra_a_formule(ws, ws_cal)

    if avvisa and abs(fattore_file - fattore) > 1e-9:
        avvisa(
            f"Il riepilogo annuale usa gia' il fattore di conversione {fattore_file:g}, "
            f"mentre questa settimana usa {fattore:g}.\n"
            f"Per coerenza con i giorni gia' presenti, le concentrazioni del riepilogo annuale "
            f"(fogli '{ws.title}' e 'Calendario') usano {fattore_file:g}: possono quindi "
            f"differire da quelle del bollettino di questa settimana.\n"
            f"Per cambiare il fattore di tutto il riepilogo basta modificare la cella "
            f"{_rif_fattore().replace('$', '')} del foglio '{ws.title}' (in alto a destra della "
            f"tabella delle concentrazioni): tutte le concentrazioni si aggiornano da sole.")

    scelta_duplicati = None
    giorni_scritti = 0

    for giorno_num in sorted(dati_settimana.keys()):
        data_str = settimana.data_giorno(giorno_num).strftime("%d/%m/%Y")
        dati = dati_settimana[giorno_num]

        riga_esistente = trova_riga_per_data(ws, data_str)
        if riga_esistente:
            if scelta_duplicati is None:
                scelta_duplicati = scegli_duplicati(data_str) if scegli_duplicati else "a"
            if scelta_duplicati == "annulla":
                continue  # giorno gia' presente, l'utente ha annullato: non lo tocca
            if scelta_duplicati == "a":
                riga_dati = scrivi_riga_annuale(ws, riga_esistente, data_str, dati, "sovrascrivi")
            elif scelta_duplicati == "b":
                riga_dati = scrivi_riga_annuale(ws, _prossima_riga_annuale(ws), data_str, dati, "nuovo")
            else:
                riga_dati = scrivi_riga_annuale(ws, riga_esistente, data_str, dati, "somma")
        else:
            riga_dati = scrivi_riga_annuale(ws, _prossima_riga_annuale(ws), data_str, dati, "nuovo")

        col_esistente = trova_colonna_per_data_calendario(ws_cal, data_str)
        if col_esistente and scelta_duplicati != "b":
            col_cal = col_esistente
        else:
            col_cal = _prossima_colonna_calendario(ws_cal)
        scrivi_colonna_calendario(ws_cal, col_cal, data_str, ws.title, riga_dati)

        giorni_scritti += 1

    crea_foglio_settimana_annuale(wb_ann, settimana.conteggi, settimana.lunedi, fattore)

    # openpyxl non salva i risultati delle formule: Excel/LibreOffice li
    # calcolano all'apertura.
    wb_ann.calculation.fullCalcOnLoad = True
    wb_ann.save(percorso)
    wb_ann.close()
    return percorso, giorni_scritti


# ============================================================
# Bollettino pollinico (.docx)
# ============================================================
def genera_bollettini_word(settimana, cartella, soglie=None):
    """Genera i bollettini Word (ITA e ENG). Ritorna la lista dei percorsi
    creati (puo' essere vuota se python-docx non e' installato o non
    ci sono dati). soglie: dict da carica_soglie(); se None viene caricato."""
    if _docx_module is None:
        return []

    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import RGBColor, Pt

    lunedi = settimana.lunedi
    lunedi_str = lunedi.strftime("%d-%m-%Y")
    fattore = settimana.fattore
    if soglie is None:
        soglie = carica_soglie(cartella) or {}

    righe_dati = dominio.righe_bollettino(settimana.conteggi, fattore, soglie)
    if not righe_dati:
        return []

    def _colore(livello):
        return dominio.LIVELLO_COLORE_WORD[livello]

    _COL_WIDTHS = [2000, 1300, 1300, 1300, 1300, 1300, 1300, 1300, 1800, 2498]

    def _set_table_widths(table):
        tbl = table._tbl
        tblGrid = tbl.find(qn("w:tblGrid"))
        if tblGrid is not None:
            for i, gc in enumerate(tblGrid.findall(qn("w:gridCol"))):
                if i < len(_COL_WIDTHS):
                    gc.set(qn("w:w"), str(_COL_WIDTHS[i]))
        tblPr = tbl.find(qn("w:tblPr"))
        if tblPr is not None:
            tblW = tblPr.find(qn("w:tblW"))
            if tblW is not None:
                tblW.set(qn("w:w"), str(sum(_COL_WIDTHS)))
                tblW.set(qn("w:type"), "dxa")

    def _set_cell_color(cell, rgb_hex):
        tc = cell._tc
        tcPr = tc.find(qn("w:tcPr"))
        if tcPr is None:
            tcPr = OxmlElement("w:tcPr")
            tc.insert(0, tcPr)
        shd = tcPr.find(qn("w:shd"))
        if shd is None:
            shd = OxmlElement("w:shd")
            tcPr.append(shd)
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), rgb_hex)

    def _set_cell_borders(cell):
        tc = cell._tc
        tcPr = tc.find(qn("w:tcPr"))
        if tcPr is None:
            tcPr = OxmlElement("w:tcPr")
            tc.insert(0, tcPr)
        tcBorders = tcPr.find(qn("w:tcBorders"))
        if tcBorders is None:
            tcBorders = OxmlElement("w:tcBorders")
            tcPr.append(tcBorders)
        for side in ("top", "left", "bottom", "right"):
            border = OxmlElement(f"w:{side}")
            border.set(qn("w:val"), "single")
            border.set(qn("w:color"), "000000")
            border.set(qn("w:sz"), "4")
            tcBorders.append(border)

    def _set_cell_paragraph_format(cell):
        tc = cell._tc
        tcPr = tc.find(qn("w:tcPr"))
        if tcPr is None:
            tcPr = OxmlElement("w:tcPr")
            tc.insert(0, tcPr)
        vAlign = tcPr.find(qn("w:vAlign"))
        if vAlign is None:
            vAlign = OxmlElement("w:vAlign")
            tcPr.append(vAlign)
        vAlign.set(qn("w:val"), "center")
        para = cell.paragraphs[0]
        pPr = para._p.find(qn("w:pPr"))
        if pPr is None:
            pPr = OxmlElement("w:pPr")
            para._p.insert(0, pPr)
        jc = pPr.find(qn("w:jc"))
        if jc is None:
            jc = OxmlElement("w:jc")
            pPr.append(jc)
        jc.set(qn("w:val"), "center")

    def _set_row_height(row, height_twips):
        trPr = row._tr.find(qn("w:trPr"))
        if trPr is None:
            trPr = OxmlElement("w:trPr")
            row._tr.insert(0, trPr)
        trH = trPr.find(qn("w:trHeight"))
        if trH is None:
            trH = OxmlElement("w:trHeight")
            trPr.append(trH)
        trH.set(qn("w:val"), str(height_twips))
        trH.set(qn("w:hRule"), "exact")

    def _set_cell_text(cell, text, italic=False, color=None, bold=None, size_pt=None):
        para = cell.paragraphs[0]
        for r in para._p.findall(qn("w:r")):
            para._p.remove(r)
        run = para.add_run(text)
        run.font.name = "Arial Narrow"
        run.italic = italic
        if color:
            run.font.color.rgb = RGBColor(int(color[:2], 16), int(color[2:4], 16), int(color[4:], 16))
        if bold is not None:
            run.font.bold = bold
        if size_pt is not None:
            run.font.size = Pt(size_pt)

    percorsi_creati = []
    for lang in ("ITA", "ENG"):
        template_name = f"{lang}_Template_Bollettino_pubblicazione.docx"
        template_path = percorsi.SCRIPT_DIR / template_name
        if not template_path.exists():
            continue

        doc = Document(str(template_path))
        table = doc.tables[0]
        tbl = table._tbl
        _set_table_widths(table)

        header_cells = table.rows[0].cells
        if lang == "ITA":
            _set_cell_text(header_cells[0], f"POLLINI – {dominio.MESI_ITA[lunedi.month - 1]} {lunedi.year}",
                           color="002060", bold=True, size_pt=10)
            giorni_long = dominio.GIORNI_ITA_LONG
        else:
            _set_cell_text(header_cells[0], f"POLLEN – {lunedi.strftime('%B')} {lunedi.year}",
                           color="002060", bold=True, size_pt=10)
            giorni_long = dominio.GIORNI_ENG_LONG

        for i in range(7):
            giorno_dt = lunedi + timedelta(days=i)
            _set_cell_text(header_cells[i + 1], f"{giorni_long[i]} {giorno_dt.day}",
                           color="002060", bold=True, size_pt=9)

        for row in list(table.rows[1:]):
            tbl.remove(row._tr)

        for riga in righe_dati:
            nome = riga.nome_ita if lang == "ITA" else riga.nome_eng
            new_row = table.add_row()
            cells = new_row.cells

            for cell in cells:
                _set_cell_borders(cell)
                _set_cell_paragraph_format(cell)

            _set_cell_text(cells[0], nome, italic=True, color="002060", size_pt=9)

            for i, liv in enumerate(riga.livelli_giorno):
                _set_cell_color(cells[i + 1], _colore(liv))

            _set_cell_color(cells[8], _colore(riga.livello_media))

            _set_row_height(new_row, 360)

        output_name = f"Bollettino_{lang}_{lunedi_str}.docx"
        output_path = Path(cartella) / output_name
        doc.save(str(output_path))
        percorsi_creati.append(output_path)

    return percorsi_creati


# ============================================================
# Cheatsheet codici specie (.docx) — foglio stampabile di riferimento
# ============================================================
def genera_cheatsheet_codici(cartella):
    """Genera un foglio stampabile (.docx) con l'elenco codice -> specie.

    Non dipende da una sessione: usa solo dominio.CODICI_SPECIE (dati
    statici), quindi si puo' generare in qualsiasi momento. Ritorna il
    percorso creato, o None se python-docx non e' installato."""
    if _docx_module is None:
        return None

    from docx import Document
    from docx.shared import Cm, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(1.2)
        section.bottom_margin = Cm(1.2)
        section.left_margin = Cm(1.5)
        section.right_margin = Cm(1.5)

    titolo = doc.add_paragraph()
    titolo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = titolo.add_run("CODICI SPECIE — CONTA POLLINICA")
    run.bold = True
    run.font.size = Pt(14)

    voci = [(codice, dominio.CODICI_SPECIE[codice]) for codice in dominio.TUTTI_CODICI]
    meta = (len(voci) + 1) // 2
    col1, col2 = voci[:meta], voci[meta:]

    table = doc.add_table(rows=meta + 1, cols=4)
    table.style = "Table Grid"
    table.autofit = False

    for i, testo in enumerate(("Cod.", "Specie", "Cod.", "Specie")):
        cell = table.cell(0, i)
        cell.text = testo
        run = cell.paragraphs[0].runs[0]
        run.bold = True
        run.font.size = Pt(10)

    for riga in range(meta):
        codice1, nome1 = col1[riga]
        table.cell(riga + 1, 0).text = codice1
        table.cell(riga + 1, 1).text = nome1
        if riga < len(col2):
            codice2, nome2 = col2[riga]
            table.cell(riga + 1, 2).text = codice2
            table.cell(riga + 1, 3).text = nome2
        for col in range(4):
            for para in table.cell(riga + 1, col).paragraphs:
                for run in para.runs:
                    run.font.size = Pt(10)

    larghezze = [Cm(1.3), Cm(6.2), Cm(1.3), Cm(6.2)]
    for i, larghezza in enumerate(larghezze):
        for cell in table.columns[i].cells:
            cell.width = larghezza

    percorso = Path(cartella) / "Cheatsheet_Codici_Specie.docx"
    percorso.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(percorso))
    return percorso
