import tempfile
import unittest
from datetime import datetime
from pathlib import Path

import openpyxl

import dominio
import esportatori
import sessione

LUNEDI = datetime(2026, 2, 9)


def _settimana_di_prova():
    s = sessione.Settimana(LUNEDI, fattore=0.4)
    s.attiva_giorno(1)
    s.inserisci(1, "01", 5)     # ACERACEAE
    s.inserisci(1, "04", 3)     # Alnus (fa parte di BETULACEAE)
    s.attiva_giorno(3)
    s.inserisci(3, "48", 10)    # Alternaria (spora)
    return s


class TestEsportaXlsx(unittest.TestCase):
    def test_scrive_conteggi_nelle_celle_giuste(self):
        s = _settimana_di_prova()
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "prova.xlsx"
            esportatori.esporta_xlsx(s, path)
            wb = openpyxl.load_workbook(path)
            ws = wb["riepilogo_settimana"]
            # codice 01 -> riga 6, lunedi -> colonna G(7)
            self.assertEqual(ws.cell(row=6, column=7).value, 5)
            # codice 04 -> riga 9
            self.assertEqual(ws.cell(row=9, column=7).value, 3)
            # codice 48 -> riga 58, mercoledi -> colonna I(9)
            self.assertEqual(ws.cell(row=58, column=9).value, 10)
            self.assertEqual(ws["Q3"].value, 0.4)
            self.assertEqual(ws["M3"].value, 2026)
            wb.close()

    def test_fattore_modificato_va_nel_file_e_nel_bollettino(self):
        s = _settimana_di_prova()
        s.imposta_fattore(0.5)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "prova.xlsx"
            esportatori.esporta_xlsx(s, path)
            wb = openpyxl.load_workbook(path)
            self.assertEqual(wb["riepilogo_settimana"]["Q3"].value, 0.5)
            wb.close()
            self.assertEqual(sessione.carica_da_xlsx(path).fattore, 0.5)
        r04 = dominio.righe_bollettino(s.conteggi, s.fattore, {})
        r04_default = dominio.righe_bollettino(s.conteggi, 0.4, {})
        self.assertNotEqual([r.concentrazioni for r in r04],
                            [r.concentrazioni for r in r04_default])

    def test_scrive_il_log(self):
        s = _settimana_di_prova()
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "prova.xlsx"
            esportatori.esporta_xlsx(s, path)
            wb = openpyxl.load_workbook(path)
            ws_log = wb["dati_grezzi"]
            n = sum(1 for row in ws_log.iter_rows(min_row=2, max_col=1) if row[0].value is not None)
            self.assertEqual(n, len(s.log))
            wb.close()

    def test_round_trip_import_export(self):
        """Esporta, poi ricarica con sessione.carica_da_xlsx: i conteggi devono
        tornare identici (verifica di parita' import/export)."""
        s = _settimana_di_prova()
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "prova.xlsx"
            esportatori.esporta_xlsx(s, path)
            ripreso = sessione.carica_da_xlsx(path)
            self.assertEqual(ripreso.conteggi, s.conteggi)
            self.assertEqual(ripreso.fattore, s.fattore)


class TestCaricaSoglie(unittest.TestCase):
    def test_trova_file_nella_cartella_dello_script(self):
        soglie = esportatori.carica_soglie()
        self.assertIsNotNone(soglie)
        self.assertIn("Alternaria", soglie)
        # unica fonte: coerente col file, non col vecchio fallback nel codice
        self.assertEqual(soglie["Cladosporium"][0], 99.9)

    def test_none_se_non_trovato(self):
        with tempfile.TemporaryDirectory() as tmp:
            import percorsi
            vecchio_script_dir = percorsi.SCRIPT_DIR
            vecchio_bundle_dir = percorsi.BUNDLE_DIR
            try:
                percorsi.SCRIPT_DIR = Path(tmp)
                percorsi.BUNDLE_DIR = Path(tmp)
                esportatori.percorsi.SCRIPT_DIR = Path(tmp)
                esportatori.percorsi.BUNDLE_DIR = Path(tmp)
                self.assertIsNone(esportatori.carica_soglie(output_dir=tmp))
            finally:
                percorsi.SCRIPT_DIR = vecchio_script_dir
                percorsi.BUNDLE_DIR = vecchio_bundle_dir
                esportatori.percorsi.SCRIPT_DIR = vecchio_script_dir
                esportatori.percorsi.BUNDLE_DIR = vecchio_bundle_dir


class TestRiepilogoAnnuale(unittest.TestCase):
    def test_crea_file_con_foglio_settimanale_e_calendario(self):
        s = _settimana_di_prova()
        with tempfile.TemporaryDirectory() as tmp:
            percorso, n = esportatori.esporta_riepilogo_annuale(s, tmp)
            self.assertIsNotNone(percorso)
            self.assertEqual(n, 2)  # 2 giorni con dati (lunedi e mercoledi)
            wb = openpyxl.load_workbook(percorso)
            self.assertIn("Calendario", wb.sheetnames)
            self.assertIn("W07", wb.sheetnames)  # 9 feb 2026 e' nella settimana ISO 7
            wb.close()

    def test_nessun_dato_non_crea_file(self):
        s = sessione.Settimana(LUNEDI)
        with tempfile.TemporaryDirectory() as tmp:
            percorso, n = esportatori.esporta_riepilogo_annuale(s, tmp)
            self.assertIsNone(percorso)
            self.assertEqual(n, 0)

    def test_giorno_duplicato_somma_se_richiesto(self):
        s = _settimana_di_prova()
        with tempfile.TemporaryDirectory() as tmp:
            esportatori.esporta_riepilogo_annuale(s, tmp)
            # riesporta la stessa settimana: il lunedi' e' gia' presente
            esportatori.esporta_riepilogo_annuale(s, tmp, scegli_duplicati=lambda data: "c")
            percorso = Path(tmp) / "Riepilogo_Annuale_2026.xlsx"
            wb = openpyxl.load_workbook(percorso)
            ws = wb.active
            riga = esportatori.trova_riga_per_data(ws, "09/02/2026")
            col = esportatori._ann_col_grezzo("01")
            self.assertEqual(ws.cell(row=riga, column=col).value, 10)  # 5 + 5
            wb.close()

    def test_giorno_duplicato_annullato_non_modifica_il_dato_esistente(self):
        s = _settimana_di_prova()
        with tempfile.TemporaryDirectory() as tmp:
            esportatori.esporta_riepilogo_annuale(s, tmp)
            # riesporta la stessa settimana ma l'utente annulla: il dato gia'
            # presente non deve essere toccato e il giorno non va contato
            percorso, n = esportatori.esporta_riepilogo_annuale(
                s, tmp, scegli_duplicati=lambda data: "annulla")
            self.assertEqual(n, 0)
            wb = openpyxl.load_workbook(percorso)
            ws = wb.active
            riga = esportatori.trova_riga_per_data(ws, "09/02/2026")
            col = esportatori._ann_col_grezzo("01")
            self.assertEqual(ws.cell(row=riga, column=col).value, 5)  # invariato
            wb.close()


class TestFattoreEFormuleAnnuale(unittest.TestCase):
    """Le concentrazioni del riepilogo annuale sono formule conta x cella del
    fattore (riferimento assoluto), come nel template settimanale."""

    def _esporta(self, tmp, **kw):
        s = kw.pop("settimana", None) or _settimana_di_prova()
        percorso, _ = esportatori.esporta_riepilogo_annuale(s, tmp, **kw)
        return percorso

    def test_cella_del_fattore_nell_angolo_in_alto_a_destra(self):
        with tempfile.TemporaryDirectory() as tmp:
            wb = openpyxl.load_workbook(self._esporta(tmp))
            ws = wb["Dati 2026"]
            self.assertEqual(ws["DS1"].value, 0.4)
            self.assertEqual(ws["DR1"].value, "Fatt.conv")
            wb.close()

    def test_concentrazioni_dati_sono_formule_con_fattore_bloccato(self):
        with tempfile.TemporaryDirectory() as tmp:
            wb = openpyxl.load_workbook(self._esporta(tmp))
            ws = wb["Dati 2026"]
            riga = esportatori.trova_riga_per_data(ws, "09/02/2026")
            lettera = openpyxl.utils.get_column_letter(esportatori._ann_col_grezzo("01"))
            conc = ws.cell(row=riga, column=esportatori._ann_col_conc("01")).value
            self.assertEqual(conc, f'=IF({lettera}{riga}="","",{lettera}{riga}*$DS$1)')
            wb.close()

    def test_calendario_rimanda_al_foglio_dati(self):
        with tempfile.TemporaryDirectory() as tmp:
            wb = openpyxl.load_workbook(self._esporta(tmp))
            cal = wb["Calendario"]
            col = esportatori.trova_colonna_per_data_calendario(cal, "09/02/2026")
            lettera = openpyxl.utils.get_column_letter(esportatori._ann_col_conc("01"))
            self.assertEqual(cal.cell(row=esportatori._cal_row_for_codice("01"), column=col).value,
                             f"='Dati 2026'!{lettera}4")
            wb.close()

    def test_foglio_settimanale_ha_il_proprio_fattore(self):
        s = sessione.Settimana(LUNEDI, fattore=0.5)
        s.attiva_giorno(1)
        s.inserisci(1, "01", 5)
        with tempfile.TemporaryDirectory() as tmp:
            wb = openpyxl.load_workbook(self._esporta(tmp, settimana=s))
            w = wb["W07"]
            riga_titolo = next(r[0].row for r in w.iter_rows() if r[0].value == "CONCENTRAZIONI (p/m3)")
            self.assertEqual(w.cell(row=riga_titolo, column=11).value, "Fatt.conv")
            self.assertEqual(w.cell(row=riga_titolo, column=12).value, 0.5)
            formule = [c.value for r in w.iter_rows(min_row=riga_titolo + 1)
                       for c in r[1:2] if isinstance(c.value, str) and c.value.startswith("=IF(")]
            self.assertTrue(formule)
            self.assertTrue(all(f"$L${riga_titolo}" in f for f in formule))
            wb.close()

    def test_somma_su_duplicato_aggiorna_la_conta_e_non_serve_ricalcolare(self):
        with tempfile.TemporaryDirectory() as tmp:
            self._esporta(tmp)
            percorso = self._esporta(tmp, scegli_duplicati=lambda data: "c")
            wb = openpyxl.load_workbook(percorso)
            ws = wb["Dati 2026"]
            riga = esportatori.trova_riga_per_data(ws, "09/02/2026")
            self.assertEqual(ws.cell(row=riga, column=esportatori._ann_col_grezzo("01")).value, 10)
            self.assertTrue(str(ws.cell(row=riga, column=esportatori._ann_col_conc("01")).value).startswith("=IF("))
            wb.close()

    def test_fattore_diverso_avvisa_e_non_sovrascrive(self):
        avvisi = []
        with tempfile.TemporaryDirectory() as tmp:
            self._esporta(tmp)
            s2 = sessione.Settimana(datetime(2026, 2, 16), fattore=0.5)
            s2.attiva_giorno(1)
            s2.inserisci(1, "01", 5)
            percorso = self._esporta(tmp, settimana=s2, avvisa=avvisi.append)
            wb = openpyxl.load_workbook(percorso)
            self.assertEqual(wb["Dati 2026"]["DS1"].value, 0.4)
            wb.close()
        self.assertEqual(len(avvisi), 1)
        self.assertIn("0.4", avvisi[0])
        self.assertIn("DS1", avvisi[0])

    def test_stesso_fattore_non_avvisa(self):
        avvisi = []
        with tempfile.TemporaryDirectory() as tmp:
            self._esporta(tmp, avvisa=avvisi.append)
            self._esporta(tmp, avvisa=avvisi.append)
        self.assertEqual(avvisi, [])

    def _rendi_file_vecchio(self, percorso, fattore):
        """Riporta il file al formato precedente: concentrazioni numeriche
        fisse e nessuna cella del fattore."""
        wb = openpyxl.load_workbook(percorso)
        ws, cal = wb["Dati 2026"], wb["Calendario"]
        for riga in range(4, ws.max_row + 1):
            for codice in dominio.TUTTI_CODICI:
                g = ws.cell(row=riga, column=esportatori._ann_col_grezzo(codice)).value
                ws.cell(row=riga, column=esportatori._ann_col_conc(codice)).value = \
                    round(g * fattore, 1) if g else None
        for col in range(2, cal.max_column + 1):
            for codice in dominio.TUTTI_CODICI:
                cal.cell(row=esportatori._cal_row_for_codice(codice), column=col).value = None
        ws["DS1"].value = None
        ws["DR1"].value = None
        wb.save(percorso)
        wb.close()

    def test_file_nel_formato_vecchio_viene_migrato_alle_formule(self):
        with tempfile.TemporaryDirectory() as tmp:
            percorso = self._esporta(tmp)
            self._rendi_file_vecchio(percorso, 0.4)
            avvisi = []
            s2 = sessione.Settimana(datetime(2026, 2, 16), fattore=0.4)
            s2.attiva_giorno(1)
            s2.inserisci(1, "01", 2)
            self._esporta(tmp, settimana=s2, avvisa=avvisi.append)
            wb = openpyxl.load_workbook(percorso)
            ws = wb["Dati 2026"]
            self.assertEqual(ws["DS1"].value, 0.4)
            riga = esportatori.trova_riga_per_data(ws, "09/02/2026")
            self.assertTrue(str(ws.cell(row=riga, column=esportatori._ann_col_conc("01")).value).startswith("=IF("))
            cal = wb["Calendario"]
            col = esportatori.trova_colonna_per_data_calendario(cal, "09/02/2026")
            self.assertTrue(str(cal.cell(row=esportatori._cal_row_for_codice("01"), column=col).value)
                            .startswith("='Dati 2026'!"))
            wb.close()
        self.assertEqual(avvisi, [])

    def test_file_vecchio_con_altro_fattore_lo_mantiene_e_avvisa(self):
        with tempfile.TemporaryDirectory() as tmp:
            percorso = self._esporta(tmp)
            self._rendi_file_vecchio(percorso, 0.4)
            avvisi = []
            s2 = sessione.Settimana(datetime(2026, 2, 16), fattore=0.5)
            s2.attiva_giorno(1)
            s2.inserisci(1, "01", 2)
            self._esporta(tmp, settimana=s2, avvisa=avvisi.append)
            wb = openpyxl.load_workbook(percorso)
            self.assertEqual(wb["Dati 2026"]["DS1"].value, 0.4)      # ricavato dai dati vecchi
            wb.close()
        self.assertEqual(len(avvisi), 1)


try:
    import docx  # noqa: F401
    _DOCX_DISPONIBILE = True
except ImportError:
    _DOCX_DISPONIBILE = False


@unittest.skipUnless(_DOCX_DISPONIBILE, "python-docx non installato")
class TestBollettiniWord(unittest.TestCase):
    def test_genera_ita_ed_eng(self):
        s = _settimana_di_prova()
        soglie = esportatori.carica_soglie() or {}
        with tempfile.TemporaryDirectory() as tmp:
            creati = esportatori.genera_bollettini_word(s, tmp, soglie=soglie)
            self.assertEqual(len(creati), 2)
            for p in creati:
                self.assertTrue(p.exists())

    def test_nessun_dato_non_genera_nulla(self):
        s = sessione.Settimana(LUNEDI)
        with tempfile.TemporaryDirectory() as tmp:
            creati = esportatori.genera_bollettini_word(s, tmp, soglie={})
            self.assertEqual(creati, [])


if __name__ == "__main__":
    unittest.main()
