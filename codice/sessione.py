#!/usr/bin/env python3
"""
Modello di sessione (Settimana) e persistenza incrementale (Journal).

Lo stato vive in memoria in un oggetto Settimana, e ogni operazione
(inserimento, annullo, correzione, nota) viene scritta subito come una riga
in un journal JSONL append-only. Il salvataggio su .xlsx avviene solo quando
l'utente lo chiede esplicitamente (comando 's'/'q' o pulsante Salva), tramite
esportatori.esporta_xlsx().

Nessuna interazione con l'utente (input/print/dialoghi): quella resta a
carico di polline_counter.py (CLI) e polline_counter_gui.py (GUI).
"""

import functools
import json
import os
from datetime import datetime, timedelta
from pathlib import Path

try:
    import openpyxl
except ImportError:
    openpyxl = None

import dominio


# ============================================================
# Modello in memoria
# ============================================================
class Settimana:
    """Stato di una settimana di conta pollinica: conteggi, log, storico."""

    def __init__(self, lunedi, fattore=dominio.FATTORE_DEFAULT):
        self.lunedi = lunedi
        self.fattore = fattore
        self.conteggi = {c: [0] * 7 for c in dominio.TUTTI_CODICI}
        self.log = []          # righe per 'dati_grezzi': dict(data,codice,specie,ora,nota)
        self.storico = []      # ultimi inserimenti: (data_str, codice, specie, quantita, ora)
        self.nome_origine = None   # nome del file/sessione da cui e' stata ripresa (o None)
        self.beep = False
        self.giorno_attivo = None
        self.ultimo_codice = None
        self._undo_stack = []
        self._listeners = []

    # ── notifiche (per aggiornare la UI a ogni cambiamento) ──
    def on_change(self, callback):
        self._listeners.append(callback)

    def _notifica(self):
        for cb in list(self._listeners):
            cb(self)

    # ── date/giorni ──
    def data_giorno(self, giorno_num):
        return self.lunedi + timedelta(days=giorno_num - 1)

    def data_str_giorno(self, giorno_num):
        return self.data_giorno(giorno_num).strftime("%d-%m-%Y")

    def nome_file_default(self):
        """Nome proposto per il file .xlsx: quello di origine, o uno nuovo dal lunedi'."""
        return self.nome_origine or f"Conta_Pollinica_{self.data_str_giorno(1)}.xlsx"

    def totale_giorno(self, giorno_num):
        idx = giorno_num - 1
        return sum(v[idx] for v in self.conteggi.values())

    def righe_specie(self, giorno_num=None):
        """Ritorna [(codice, specie, valore_o_lista)] per le specie con dati > 0.
        Se giorno_num e' None, valore = totale settimana; altrimenti valore del giorno."""
        righe = []
        for codice in dominio.TUTTI_CODICI:
            vals = self.conteggi[codice]
            if giorno_num is None:
                tot = sum(vals)
                if tot > 0:
                    righe.append((codice, dominio.CODICI_SPECIE[codice], tot))
            else:
                v = vals[giorno_num - 1]
                if v > 0:
                    righe.append((codice, dominio.CODICI_SPECIE[codice], v))
        return righe

    # ── operazioni mutanti ──
    def _registra_log(self, giorno_num, codice, specie, ora, nota):
        """Aggiunge una riga a ``log`` (ora = adesso se non data). Ritorna il suo indice."""
        self.log.append({"data": self.data_str_giorno(giorno_num), "codice": codice,
                         "specie": specie, "ora": ora or datetime.now().strftime("%H:%M:%S"),
                         "nota": nota})
        return len(self.log) - 1

    def attiva_giorno(self, giorno_num, journal=None):
        """Rende 'giorno_num' il giorno corrente: azzera lo storico di annullo
        (identico al comportamento originale: 'u' funziona solo nella sessione
        del giorno in corso)."""
        self.giorno_attivo = giorno_num
        self._undo_stack = []
        if journal:
            journal.registra("attiva", giorno=giorno_num)
        self._notifica()

    def inserisci(self, giorno_num, codice, quantita, journal=None, ora=None):
        idx = giorno_num - 1
        specie = dominio.CODICI_SPECIE[codice]
        data_str = self.data_str_giorno(giorno_num)
        ora = ora or datetime.now().strftime("%H:%M:%S")

        self.conteggi[codice][idx] += quantita
        indici = [self._registra_log(giorno_num, codice, specie, ora, None)
                  for _ in range(quantita)]
        self.storico.append((data_str, codice, specie, quantita, ora))
        if self.giorno_attivo == giorno_num:
            self._undo_stack.append({"codice": codice, "quantita": quantita, "log_indici": indici})
        self.ultimo_codice = codice

        if journal:
            journal.registra("add", giorno=giorno_num, codice=codice, quantita=quantita)
        self._notifica()
        return self.conteggi[codice][idx]

    def annulla(self, giorno_num, journal=None, ora=None):
        """Annulla l'ultimo inserimento del giorno attivo (rimuove anche la voce
        dallo storico, non solo dal conteggio).
        Ritorna (codice, quantita, nuovo_totale) o None se non c'e' nulla da annullare."""
        if self.giorno_attivo != giorno_num or not self._undo_stack:
            return None
        op = self._undo_stack.pop()
        codice, quantita, indici = op["codice"], op["quantita"], op["log_indici"]
        idx = giorno_num - 1

        self.conteggi[codice][idx] = max(0, self.conteggi[codice][idx] - quantita)
        for i in sorted(indici, reverse=True):
            if 0 <= i < len(self.log):
                del self.log[i]
        if self.storico:
            ultimo = self.storico[-1]
            if (ultimo[1] == codice and ultimo[3] == quantita
                    and ultimo[0] == self.data_str_giorno(giorno_num)):
                self.storico.pop()

        self._registra_log(giorno_num, codice, dominio.CODICI_SPECIE[codice], ora,
                           f"ANNULLATO x{quantita}")

        if journal:
            journal.registra("undo", giorno=giorno_num)
        self._notifica()
        return codice, quantita, self.conteggi[codice][idx]

    def correggi(self, giorno_num, codice, nuovo_valore, journal=None, ora=None):
        if nuovo_valore < 0:
            raise ValueError("Il valore non puo' essere negativo.")
        idx = giorno_num - 1
        vecchio = self.conteggi[codice][idx]
        self.conteggi[codice][idx] = nuovo_valore

        self._registra_log(giorno_num, codice, dominio.CODICI_SPECIE[codice], ora,
                           f"CORREZIONE {vecchio}->{nuovo_valore}")

        if journal:
            journal.registra("correggi", giorno=giorno_num, codice=codice, valore=nuovo_valore)
        self._notifica()
        return vecchio

    def aggiungi_nota(self, giorno_num, testo, journal=None, ora=None):
        self._registra_log(giorno_num, "--", "NOTA", ora, testo)
        if journal:
            journal.registra("nota", giorno=giorno_num, testo=testo)
        self._notifica()


# ============================================================
# Journal: persistenza incrementale su disco (JSONL)
# ============================================================
class Journal:
    """Log append-only della sessione corrente.

    La prima riga ('inizio') e' uno snapshot completo dello stato di
    partenza (conteggi, log, fattore); le righe successive sono eventi
    incrementali. In caso di crash, ripristina_da_journal() ricostruisce
    lo stato rigiocando esattamente le stesse operazioni.
    """

    def __init__(self, path):
        self.path = Path(path)
        self._file = None

    def avvia(self, settimana):
        """Apre il file (sovrascrivendo un eventuale journal precedente con lo
        stesso nome) e registra lo snapshot iniziale."""
        self._file = open(self.path, "w", encoding="utf-8")
        self._scrivi({
            "t": "inizio",
            "lunedi": settimana.lunedi.strftime("%Y-%m-%d"),
            "fattore": settimana.fattore,
            "conteggi": settimana.conteggi,
            "log": settimana.log,
            "origine": settimana.nome_origine,
        })

    @classmethod
    def riprendi(cls, path):
        """Riapre in append un journal esistente (dopo un ripristino da crash):
        NON riscrive lo snapshot iniziale, gia' presente nel file."""
        j = cls(path)
        j._file = open(j.path, "a", encoding="utf-8")
        return j

    def registra(self, tipo, **campi):
        if self._file is None:
            return
        riga = {"t": tipo, "ts": datetime.now().strftime("%H:%M:%S")}
        riga.update(campi)
        self._scrivi(riga)

    def _scrivi(self, riga):
        self._file.write(json.dumps(riga, ensure_ascii=False) + "\n")
        self._file.flush()
        try:
            os.fsync(self._file.fileno())
        except OSError:
            pass

    def chiudi(self):
        if self._file:
            self._file.close()
            self._file = None

    def elimina(self):
        """Chiude e cancella il file di journal (dopo un salvataggio riuscito)."""
        self.chiudi()
        try:
            self.path.unlink(missing_ok=True)
        except Exception:
            pass


def nome_journal(lunedi):
    return f"~sessione_{lunedi.strftime('%d-%m-%Y')}.jsonl"


def ripristina_da_journal(path):
    """Ricostruisce una Settimana rigiocando integralmente un file journal."""
    righe = []
    with open(path, "r", encoding="utf-8") as fh:
        for linea in fh:
            linea = linea.strip()
            if linea:
                righe.append(json.loads(linea))

    if not righe or righe[0].get("t") != "inizio":
        raise ValueError(f"Il file di sessione '{path}' e' danneggiato o vuoto.")

    inizio = righe[0]
    lunedi = datetime.strptime(inizio["lunedi"], "%Y-%m-%d")
    settimana = Settimana(lunedi, fattore=inizio.get("fattore", dominio.FATTORE_DEFAULT))
    settimana.nome_origine = inizio.get("origine")
    for codice, valori in (inizio.get("conteggi") or {}).items():
        if codice in settimana.conteggi:
            settimana.conteggi[codice] = list(valori)
    settimana.log = list(inizio.get("log") or [])

    for riga in righe[1:]:
        tipo = riga.get("t")
        ora = riga.get("ts")
        if tipo == "attiva":
            settimana.attiva_giorno(riga["giorno"])
        elif tipo == "add":
            settimana.inserisci(riga["giorno"], riga["codice"], riga["quantita"], ora=ora)
        elif tipo == "undo":
            settimana.annulla(riga["giorno"], ora=ora)
        elif tipo == "correggi":
            settimana.correggi(riga["giorno"], riga["codice"], riga["valore"], ora=ora)
        elif tipo == "nota":
            settimana.aggiungi_nota(riga["giorno"], riga["testo"], ora=ora)
    return settimana


def recupera_sessioni(cartella):
    """Elenca i journal recuperabili in 'cartella': [(path, info_dict), ...],
    piu' recenti prima. info_dict ha 'lunedi', 'origine', 'n_operazioni'."""
    risultati = []
    for f in sorted(Path(cartella).glob("~sessione_*.jsonl"),
                     key=lambda p: p.stat().st_mtime, reverse=True):
        info = _leggi_intestazione_journal(f)
        if info:
            risultati.append((f, info))
    return risultati


def _leggi_intestazione_journal(path):
    try:
        with open(path, "r", encoding="utf-8") as fh:
            righe = fh.readlines()
        if not righe:
            return None
        prima = json.loads(righe[0])
        if prima.get("t") != "inizio":
            return None
        return {
            "lunedi": prima["lunedi"],
            "origine": prima.get("origine"),
            "n_operazioni": max(0, len(righe) - 1),
        }
    except Exception:
        return None


# ============================================================
# Import da un file .xlsx esistente
# ============================================================
def carica_da_xlsx(path):
    """Legge un file .xlsx esistente e ricostruisce una Settimana.

    Solleva ValueError se il file non ha la struttura attesa (es. un riepilogo
    annuale senza il foglio 'riepilogo_settimana')."""
    if openpyxl is None:
        raise RuntimeError("openpyxl non installato.")

    path = Path(path)
    wb = openpyxl.load_workbook(path)
    try:
        if "riepilogo_settimana" not in wb.sheetnames:
            raise ValueError(
                f"'{path.name}' non e' un file di conta pollinica valido "
                f"(manca il foglio 'riepilogo_settimana')."
            )
        ws = wb["riepilogo_settimana"]

        lunedi = None
        val = ws["J3"].value
        if isinstance(val, datetime):
            lunedi = dominio.lunedi_di(val)
        elif isinstance(val, str):
            dt = dominio.parse_data_flessibile(val)
            if dt:
                lunedi = dominio.lunedi_di(dt)
        if lunedi is None:
            lunedi = dominio.lunedi_di(datetime.now())

        fattore_val = ws["Q3"].value
        fattore = float(fattore_val) if isinstance(fattore_val, (int, float)) and fattore_val > 0 \
            else dominio.FATTORE_DEFAULT

        settimana = Settimana(lunedi, fattore=fattore)
        for codice in dominio.TUTTI_CODICI:
            row = dominio.codice_to_row(codice)
            for g in range(1, 8):
                col = dominio.giorno_to_col(g)
                v = ws.cell(row=row, column=col).value
                if isinstance(v, (int, float)) and v:
                    settimana.conteggi[codice][g - 1] = int(v)

        if "dati_grezzi" in wb.sheetnames:
            ws_log = wb["dati_grezzi"]
            for row in ws_log.iter_rows(min_row=2):
                if row[0].value is None:
                    continue
                data_v = row[0].value
                codice_v = row[1].value if len(row) > 1 else None
                specie_v = row[2].value if len(row) > 2 else None
                ora_v = row[3].value if len(row) > 3 else None
                nota_v = row[4].value if len(row) > 4 else None
                settimana.log.append({
                    "data": str(data_v), "codice": codice_v, "specie": specie_v,
                    "ora": str(ora_v) if ora_v else "", "nota": nota_v,
                })

        settimana.nome_origine = path.name
        return settimana
    finally:
        wb.close()


def cerca_file_ripresa(cartella, template_name):
    """Cerca file .xlsx validi per la ripresa: esclude il template e i
    riepiloghi annuali, e verifica che abbiano il foglio 'riepilogo_settimana'.
    Ritorna una lista di Path, piu' recenti prima."""
    cartella = Path(cartella)
    candidati = [
        f for f in cartella.glob("*.xlsx")
        if f.is_file() and f.name != template_name
        and not f.name.startswith("Riepilogo_Annuale_")
    ]
    validi = [f for f in candidati if _ha_foglio_riepilogo(f)]
    validi.sort(key=lambda f: f.stat().st_mtime, reverse=True)
    return validi


def _ha_foglio_riepilogo(path):
    if openpyxl is None:
        return False
    try:
        wb = openpyxl.load_workbook(path, read_only=True)
        ok = "riepilogo_settimana" in wb.sheetnames
        wb.close()
        return ok
    except Exception:
        return False


def conta_righe_log(path):
    """Stringa '(N righe log)' per un file .xlsx, o '' se non leggibile. Aprire
    il file costa ~0,25 s: il risultato si ricorda finche' il file non cambia
    (la schermata iniziale lo richiede a ogni rientro)."""
    try:
        stato = Path(path).stat()
    except OSError:
        return ""
    try:
        return _conta_righe_log(str(path), stato.st_mtime_ns, stato.st_size)
    except Exception:
        # Errore (file bloccato da un altro programma, scrittura in corso): le
        # eccezioni non entrano nella cache, al prossimo rientro si riprova.
        return ""


@functools.lru_cache(maxsize=64)
def _conta_righe_log(path, _mtime_ns, _dimensione):
    if openpyxl is None:
        return ""
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        if "dati_grezzi" not in wb.sheetnames:
            return ""
        ws = wb["dati_grezzi"]
        n = sum(1 for row in ws.iter_rows(min_row=2, max_col=1) if row[0].value is not None)
        return f"({n} righe log)"
    finally:
        wb.close()


# ============================================================
# Configurazione cartella di lavoro (pollencounter.cfg)
# ============================================================
def _leggi_config(config_file):
    try:
        return json.loads(Path(config_file).read_text(encoding="utf-8"))
    except Exception:       # file assente o illeggibile: configurazione vuota
        return {}


def _scrivi_config(config_file, config):
    Path(config_file).write_text(json.dumps(config, ensure_ascii=False, indent=2),
                                 encoding="utf-8")


def leggi_cartella_anno(config_file, anno):
    """Ritorna il Path configurato per 'anno', o None se assente o non piu' valido."""
    config = _leggi_config(config_file)
    chiave = str(anno)
    if chiave in config:
        cartella = Path(config[chiave])
        if cartella.exists():
            return cartella
    return None


def salva_cartella_anno(config_file, anno, cartella):
    config = _leggi_config(config_file)
    config[str(anno)] = str(cartella)
    _scrivi_config(config_file, config)


def leggi_sinonimi_vocali(config_file):
    """Sinonimi vocali personali da pollencounter.cfg, chiave 'sinonimi_vocali':
    {"24": ["erba", ...]}. Ritorna {} se assenti o malformati (mai un errore:
    la voce funziona comunque con i sinonimi di default)."""
    valore = _leggi_config(config_file).get("sinonimi_vocali")
    if not isinstance(valore, dict):
        return {}
    return {str(codice): [f for f in forme if isinstance(f, str)]
            for codice, forme in valore.items() if isinstance(forme, list)}


def leggi_parola_attivazione(config_file):
    """Parola di attivazione della voce da pollencounter.cfg, chiave
    'parola_attivazione'. Assente o non testuale -> default ('conta');
    stringa vuota -> prefisso disattivato."""
    valore = _leggi_config(config_file).get("parola_attivazione")
    if not isinstance(valore, str):
        return dominio.PAROLA_ATTIVAZIONE_DEFAULT
    return dominio.normalizza_parlato(valore)


DIVISIONE_FINESTRA_MIN, DIVISIONE_FINESTRA_MAX = 0.2, 0.8


def leggi_layout_finestra(config_file):
    """Dimensioni e proporzione dei pannelli salvate in pollencounter.cfg
    (chiave 'finestra': larghezza, altezza, divisione). Ritorna un dict con le
    sole voci valide (anche vuoto): chi chiama usa i propri valori predefiniti
    per quelle mancanti. Non salva la posizione: con un monitor diverso la
    finestra potrebbe finire fuori schermo."""
    valore = _leggi_config(config_file).get("finestra")
    if not isinstance(valore, dict):
        return {}
    layout = {}
    for chiave in ("larghezza", "altezza"):
        v = valore.get(chiave)
        if isinstance(v, int) and not isinstance(v, bool) and 300 <= v <= 10000:
            layout[chiave] = v
    d = valore.get("divisione")
    if (isinstance(d, (int, float)) and not isinstance(d, bool)
            and DIVISIONE_FINESTRA_MIN <= d <= DIVISIONE_FINESTRA_MAX):
        layout["divisione"] = float(d)
    return layout


def salva_layout_finestra(config_file, larghezza, altezza, divisione):
    """Salva le dimensioni della finestra senza toccare le altre chiavi. Valori
    non sensati (finestra ridotta a icona, divisione fuori limiti) si ignorano."""
    config = _leggi_config(config_file)
    layout = {}
    if 300 <= larghezza <= 10000 and 300 <= altezza <= 10000:
        layout["larghezza"], layout["altezza"] = int(larghezza), int(altezza)
    if DIVISIONE_FINESTRA_MIN <= divisione <= DIVISIONE_FINESTRA_MAX:
        layout["divisione"] = round(float(divisione), 3)
    if not layout:
        return
    config["finestra"] = layout
    _scrivi_config(config_file, config)


ATTESA_VOCALE_DEFAULT_MS = 300


def leggi_attesa_vocale(config_file):
    """Millisecondi di pausa dopo i quali una frase vocale gia' valida viene
    chiusa senza aspettare il silenzio lungo del riconoscitore (chiave
    'voce_attesa_ms' di pollencounter.cfg). Assente o non valida -> 300;
    0 -> chiusura anticipata disattivata; limitata a 2000."""
    valore = _leggi_config(config_file).get("voce_attesa_ms")
    if isinstance(valore, bool) or not isinstance(valore, (int, float)) or valore < 0:
        return ATTESA_VOCALE_DEFAULT_MS
    return int(min(valore, 2000))
